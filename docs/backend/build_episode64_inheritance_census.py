#!/usr/bin/env python3
"""Build the read-only Episode 64 inheritance-census input.

The builder joins the authoritative 1,993-face workbook to facts already stored
in Backend Authority.  It never edits execution contracts and never supplies
missing semantics.  Mechanical identity states remain explicit so the LLM can
perform the later semantic census without receiving fabricated answers.
"""
from __future__ import annotations

import hashlib
import json
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DOCS = Path(__file__).resolve().parents[1]
REPOSITORY = DOCS.parent
CONTRACTS = DOCS / "backend/target_execution_contracts/contracts"
CENSUS = DOCS / "authority-evidence/platform/episode64-the13-authoritative-census-2026-10-02/SCZN3_EP64_THE13_1993_CONSOLIDATED_MASTER.xlsx"
CENSUS_SHEET = "THE 13 — 1993 Master"
PUBLISHER_SOURCE = DOCS / "authority-evidence/publisher-os/P_OS_12_PUBLISHER_TARGET_ONLY_2026-09-08.zip"
QT_AUTHORITY = DOCS / "backend/target_knowledge/qt_episode64_smart_target_authority.json"
FOUNDER_DIR = DOCS / "authority-evidence/platform/episode64-founder-walk-2026-10-02"
FOUNDER_LEDGER = FOUNDER_DIR / "FOUNDER_WALK_LEDGER.json"
FOUNDER_CLOSURE = FOUNDER_DIR / "FOUNDER_20_OF_20_UNIQUE_AUDIT_CLOSURE.json"
SCORE_ON_TOUCH = FOUNDER_DIR / "SCZN3_SCORE_ON_TOUCH_FOUNDER_AUTHORITY.json"
PRINTER_SHOOTER_AUTHORITY = DOCS / "governance/EPISODE64_FOUNDER_AUTHORITY_PRINTER_SHOOTER_SCZN3_2026-10-03.md"
X10_AUTHORITY = DOCS / "governance/EPISODE64_FOUNDER_AUTHORITY_SCZN3_X10_ORDERED_GRAMMAR_2026-10-03.md"
OUTPUT_DIR = DOCS / "authority-evidence/platform/episode64-inheritance-census-input-2026-10-04"

FULL_CENSUS = OUTPUT_DIR / "SCZN3_EP64_THE13_1993_INHERITANCE_CENSUS.json"
SUMMARY = OUTPUT_DIR / "SCZN3_EP64_THE13_1993_ACCOUNTING_SUMMARY.json"
SCHEMA = OUTPUT_DIR / "SCZN3_EP64_THE13_1993_CENSUS_SCHEMA.json"
PROVENANCE = OUTPUT_DIR / "SCZN3_EP64_THE13_1993_PROVENANCE_MANIFEST.json"
UNRECONCILED = OUTPUT_DIR / "SCZN3_EP64_THE13_1993_UNRECONCILED.json"
REPORT = OUTPUT_DIR / "SCZN3_EP64_THE13_1993_RECONCILIATION_REPORT.md"
DELIVERY = OUTPUT_DIR / "SCZN3_EP64_THE13_1993_DELIVERY_MANIFEST.json"

EXPECTED_TOTAL = 1993
EXPECTED_PRINTER_TOTALS = {
    "Baker Targets": 343,
    "GunFun": 241,
    "The National Target Company": 302,
    "Highwild": 45,
    "American Target Company": 197,
    "Speedwell Targets": 163,
    "Thompson Target Inc.": 97,
    "EZ Aim / Allen Company": 87,
    "Champion Targets / Primos": 44,
    "U.S. Target Inc.": 174,
    "Birchwood Casey": 34,
    "Splatterburst Targets": 12,
    "Qualification Targets": 254,
}

PUBLISHER_ALIASES = {
    "Qualification Targets": {"Qualification Targets Inc"},
    "GunFun": {"GunFun", "GunFun Shooting Targets, Inc."},
}

PUBLISHER_FILES = {
    "Baker Targets": "baker-targets",
    "American Target Company": "american-target",
    "The National Target Company": "national-target",
    "U.S. Target Inc.": "us-target",
    "Highwild": "highwild",
    "GunFun Shooting Targets, Inc.": "gunfun",
    "GunFun": "gunfun",
    "Speedwell Targets": "speedwell",
    "Thompson Target Inc.": "thompson-target",
    "EZ Aim / Allen Company": "ez-aim-allen",
    "Birchwood Casey": "birchwood-casey",
    "Champion Targets / Primos": "champion",
    "Splatterburst Targets": "splatterburst",
}

UNKNOWN = "UNKNOWN"
NOT_STORED = "NOT_STORED"


def clean(value: Any) -> str:
    return str(value or "").strip()


def normalized(value: Any) -> str:
    ascii_value = unicodedata.normalize("NFKD", clean(value)).encode("ascii", "ignore").decode().casefold()
    return re.sub(r"[^a-z0-9]+", " ", ascii_value).strip()


def split_tokens(value: Any) -> list[str]:
    return [normalized(item) for item in re.split(r"[,|\n]", clean(value)) if normalized(item)]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def relative(path: Path) -> str:
    return path.relative_to(REPOSITORY).as_posix()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def canonical_json(value: Any) -> str:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False)


def registration_definition(contract: dict) -> dict:
    return (((contract.get("components") or {}).get("registration") or {}).get("definition") or {})


def canonical_face_id(contract: dict) -> str | None:
    value = clean(registration_definition(contract).get("assetSha256")).lower()
    return value if re.fullmatch(r"[0-9a-f]{64}", value) else None


def component_definition(contract: dict, component: str) -> dict | None:
    item = (contract.get("components") or {}).get(component)
    return (item or {}).get("definition") if item else None


def source_keys(record: dict) -> set[str]:
    product = record.get("product") or {}
    return {
        clean(value)
        for value in (
            record.get("recordId"), record.get("sourceProductId"), record.get("sourceVariantId"),
            record.get("sku"), record.get("id"), record.get("catalogId"),
            record.get("wordpressProductId"), product.get("sku"),
        )
        if clean(value)
    }


def load_publisher_source_indexes() -> dict[str, dict[str, dict]]:
    indexes: dict[str, dict[str, dict]] = {}
    with zipfile.ZipFile(PUBLISHER_SOURCE) as archive:
        for publisher, filename in PUBLISHER_FILES.items():
            index: dict[str, dict] = {}
            for record in json.loads(archive.read(f"publishers/{filename}.json")):
                for key in source_keys(record):
                    index.setdefault(key, record)
            indexes[publisher] = index
    return indexes


def lookup_publisher_source(indexes: dict[str, dict[str, dict]], contract: dict) -> dict | None:
    target = contract.get("targetProfile") or {}
    publisher = clean((target.get("publisher") or {}).get("name"))
    product = target.get("product") or {}
    index = indexes.get(publisher, {})
    product_id = clean(product.get("id"))
    for value in (product.get("sku"), product_id, product_id.removeprefix("p-os-")):
        if clean(value) in index:
            return index[clean(value)]
    return None


def load_census() -> tuple[list[dict], dict[str, int]]:
    workbook = load_workbook(CENSUS, read_only=True, data_only=True)
    sheet = workbook[CENSUS_SHEET]
    headers = [clean(value) for value in next(sheet.iter_rows(values_only=True))]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        source = dict(zip(headers, values))
        master_number = int(source["Master #"])
        rows.append({
            "censusFaceKey": f"THE13-MASTER-{master_number:04d}",
            "masterNumber": master_number,
            "printer": clean(source["Printer"]),
            "targetName": clean(source["Target name"]),
            "recordSource": clean(source["Record source"]),
            "sourceRecordCount": int(source["Source record count"] or 0),
            "inventoryNumbersOrQtOrdinal": clean(source["Inventory #s / QT ordinal"]) or NOT_STORED,
            "inventoryRecordIdsOrQtIdentityId": clean(source["Inventory record IDs / QT Identity ID"]) or NOT_STORED,
            "printerProductIdsOrSkus": clean(source["Printer product IDs / SKU"]) or NOT_STORED,
            "assignedFamilyOrQtFamily": clean(source["Assigned family / QT SCZN3 family"]) or NOT_STORED,
            "smartFamily": clean(source["Smart Family"]) or NOT_STORED,
            "statusOrQtCapability": clean(source["Status / QT capability"]) or NOT_STORED,
            "sourceUrlsOrPublisherEvidence": clean(source["Source URLs / Publisher evidence"]) or NOT_STORED,
            "image": clean(source["Image"]) or NOT_STORED,
            "workbookAssetSha256": clean(source["SHA-256"]).lower() or NOT_STORED,
            "explicitGovernedLabels": clean(source["Explicit governed labels"]) or NOT_STORED,
            "explicitGovernedValues": clean(source["Explicit governed values"]) or NOT_STORED,
            "registeredSemanticRegionsOrObjects": clean(source["Registered semantic regions/objects"]) or NOT_STORED,
            "storedGoverningRule": clean(source["Stored governing rule"]) or NOT_STORED,
            "codyActionStored": clean(source["Cody action"]) or NOT_STORED,
            "geometryInstruction": clean(source["Geometry instruction"]) or NOT_STORED,
            "authorityLimitOrUnresolved": clean(source["Authority limit / unresolved"]) or NOT_STORED,
            "forbiddenInference": clean(source["Forbidden inference"]) or NOT_STORED,
            "requiredPersistence": clean(source["Required persistence"]) or NOT_STORED,
            "classificationBasis": clean(source["Classification basis"]) or NOT_STORED,
        })
    totals_sheet = workbook["Printer Totals — 13"]
    workbook_totals = {
        clean(row[1]): int(row[2])
        for row in totals_sheet.iter_rows(min_row=2, values_only=True)
        if clean(row[1]) and clean(row[1]) != "THE 13 TOTAL"
    }
    workbook.close()
    return rows, workbook_totals


def load_contract_inventory() -> tuple[dict[str, dict], list[dict]]:
    grouped: dict[tuple[str, str], list[tuple[Path, dict, str]]] = defaultdict(list)
    for path in sorted(CONTRACTS.glob("*.json")):
        package = read_json(path)
        identity = (clean(package.get("contractId")), clean(package.get("contractVersion")))
        grouped[identity].append((path, package, hashlib.sha256(canonical_json(package).encode()).hexdigest()))

    selected: dict[str, dict] = {}
    conflicts = []
    for identity, variants in sorted(grouped.items()):
        canonical_candidates = [item for item in variants if not re.search(r" \d+$", item[0].stem)]
        chosen = (canonical_candidates or variants)[0]
        path, package, content_digest = chosen
        runtime_id = clean((package.get("targetProfile") or {}).get("id"))
        if not runtime_id:
            runtime_id = f"MISSING_RUNTIME_ID:{identity[0]}@{identity[1]}"
        variant_digests = sorted({item[2] for item in variants})
        selected[runtime_id] = {
            "path": path,
            "package": package,
            "identity": identity,
            "contentDigest": content_digest,
            "variantPaths": [relative(item[0]) for item in variants],
            "variantContentDigests": variant_digests,
            "hasConflictingVariants": len(variant_digests) > 1,
        }
        if len(variant_digests) > 1:
            conflicts.append({
                "issue": "CONFLICTING_CONTRACT_FILE_VARIANTS",
                "contractId": identity[0],
                "contractVersion": identity[1],
                "selectedCanonicalPath": relative(path),
                "variantPaths": [relative(item[0]) for item in variants],
                "variantContentDigests": variant_digests,
            })
    return selected, conflicts


def load_qt_authority() -> tuple[dict[str, dict], dict]:
    authority = read_json(QT_AUTHORITY)
    chairs = {clean(chair.get("qtCatalogId")): chair for chair in authority.get("chairs", [])}
    header = {key: value for key, value in authority.items() if key != "chairs"}
    return chairs, header


def load_founder_history() -> tuple[dict[str, list[dict]], dict[str, list[dict]], dict]:
    by_runtime: dict[str, list[dict]] = defaultdict(list)
    by_face: dict[str, list[dict]] = defaultdict(list)
    ledger = read_json(FOUNDER_LEDGER)
    for entry in ledger.get("entries", []):
        record = {"source": relative(FOUNDER_LEDGER), **entry}
        if clean(entry.get("runtimeId")):
            by_runtime[clean(entry.get("runtimeId"))].append(record)
    for path in sorted(FOUNDER_DIR.glob("FOUNDER_*_FOUNDER_APPROVAL.json")):
        approval = read_json(path)
        record = {"source": relative(path), **approval}
        if clean(approval.get("canonicalFaceId")):
            by_face[clean(approval.get("canonicalFaceId")).lower()].append(record)
        for runtime_id in ((approval.get("target") or {}).get("runtimeIds") or []):
            by_runtime[clean(runtime_id)].append(record)
    closure = read_json(FOUNDER_CLOSURE)
    return by_runtime, by_face, closure


def contract_summary(contract: dict) -> dict:
    target = contract.get("targetProfile") or {}
    product = target.get("product") or {}
    registration = registration_definition(contract)
    geometry = component_definition(contract, "geometry") or {}
    scoring = component_definition(contract, "scoring") or {}
    mission = component_definition(contract, "mission") or {}
    components = contract.get("components") or {}
    return {
        "runtimeId": target.get("id") or UNKNOWN,
        "contractId": contract.get("contractId") or UNKNOWN,
        "contractVersion": contract.get("contractVersion") or UNKNOWN,
        "canonicalFaceId": canonical_face_id(contract) or UNKNOWN,
        "printer": (target.get("publisher") or {}).get("name") or UNKNOWN,
        "productId": product.get("id") or NOT_STORED,
        "productName": product.get("name") or NOT_STORED,
        "sku": product.get("sku") or NOT_STORED,
        "publisherUrl": (target.get("publisher") or {}).get("url") or NOT_STORED,
        "productUrl": product.get("productUrl") or product.get("sourceUrl") or NOT_STORED,
        "missionFamily": contract.get("missionFamily") or NOT_STORED,
        "resultPackageType": contract.get("resultPackageType") or NOT_STORED,
        "executor": contract.get("executor") or NOT_STORED,
        "secAdapter": contract.get("secAdapter") or NOT_STORED,
        "preservationValidator": contract.get("preservationValidator") or NOT_STORED,
        "geometryStatus": (components.get("geometry") or {}).get("status") or NOT_STORED,
        "geometryLayoutModel": geometry.get("layoutModel") or NOT_STORED,
        "governedRegionCount": len(geometry.get("regions") or []),
        "scoringStatus": (components.get("scoring") or {}).get("status") or NOT_STORED,
        "scoringModel": scoring.get("model") or NOT_STORED,
        "classificationFootprintMode": scoring.get("classificationFootprintMode") or NOT_STORED,
        "outsideRule": scoring.get("outsideRule") or NOT_STORED,
        "publisherRuleAuthority": scoring.get("publisherRuleAuthority") or NOT_STORED,
        "missionDefinitionStored": bool(mission),
        "registrationAssetPath": registration.get("assetPath") or NOT_STORED,
        "authorityContractVersion": contract.get("authorityContractVersion") or NOT_STORED,
        "approval": contract.get("approval") or NOT_STORED,
        "lifecycleStatus": contract.get("lifecycleStatus") or NOT_STORED,
        "supportStatus": contract.get("supportStatus") or NOT_STORED,
        "capabilityStates": ((contract.get("session") or {}).get("capabilityStates") or NOT_STORED),
        "unsupportedUses": contract.get("unsupportedUses") or [],
    }


def match_contracts(face: dict, inventory: dict[str, dict]) -> tuple[list[str], str, dict]:
    allowed_publishers = PUBLISHER_ALIASES.get(face["printer"], {face["printer"]})
    pool = []
    for runtime_id, item in inventory.items():
        target = item["package"].get("targetProfile") or {}
        if clean((target.get("publisher") or {}).get("name")) in allowed_publishers:
            pool.append((runtime_id, item))

    identity_tokens = [clean(item) for item in re.split(r"[,\n]", face["inventoryRecordIdsOrQtIdentityId"]) if clean(item)]
    census_tokens = set(split_tokens(face["printerProductIdsOrSkus"]))
    qt_matches = []
    token_matches = []
    name_matches = []
    for runtime_id, item in pool:
        target = item["package"].get("targetProfile") or {}
        product = target.get("product") or {}
        product_ids = {clean(product.get("id")), clean(product.get("qtCatalogId"))}
        product_ids.discard("")
        if product_ids.intersection(identity_tokens):
            qt_matches.append(runtime_id)
        contract_tokens = set(split_tokens(product.get("sku")))
        contract_tokens.add(normalized(product.get("id")))
        contract_tokens.discard("")
        if census_tokens.intersection(contract_tokens):
            token_matches.append(runtime_id)
        if normalized(product.get("name")) == normalized(face["targetName"]):
            name_matches.append(runtime_id)

    if qt_matches:
        matches, method = qt_matches, "exact_qt_identity_id"
    elif token_matches:
        matches, method = token_matches, "exact_sku_or_product_id_token"
    elif name_matches:
        matches, method = name_matches, "exact_normalized_product_name"
    else:
        matches, method = [], "NO_RUNTIME_BINDING_MATCH"
    matches = sorted(set(matches))
    return matches, method, {
        "censusIdentityTokens": identity_tokens,
        "censusProductTokens": sorted(census_tokens),
        "candidateRuntimeCount": len(matches),
    }


def distinct_non_missing(values: list[Any]) -> list[Any]:
    encoded = {}
    for value in values:
        if value in (None, "", NOT_STORED, UNKNOWN):
            continue
        encoded.setdefault(canonical_json(value), value)
    return [encoded[key] for key in sorted(encoded)]


def founder_lessons(records: list[dict]) -> list[str] | str:
    lessons = []
    for record in records:
        lessons.extend(clean(item) for item in (record.get("repairs") or []) if clean(item))
        if clean(record.get("founderCharacterization")):
            lessons.append(clean(record["founderCharacterization"]))
    return sorted(set(lessons)) or NOT_STORED


def build() -> tuple[dict, dict, dict, dict, str]:
    census_rows, workbook_totals = load_census()
    inventory, contract_conflicts = load_contract_inventory()
    publisher_indexes = load_publisher_source_indexes()
    qt_chairs, qt_authority_header = load_qt_authority()
    founder_by_runtime, founder_by_face, founder_closure = load_founder_history()

    runtime_bindings: dict[str, dict] = {}
    for runtime_id, item in sorted(inventory.items()):
        contract = item["package"]
        source_record = lookup_publisher_source(publisher_indexes, contract)
        runtime_bindings[runtime_id] = {
            "summary": contract_summary(contract),
            "selectedContractPath": relative(item["path"]),
            "selectedContractContentSha256": item["contentDigest"],
            "inventoryVariantPaths": item["variantPaths"],
            "inventoryVariantContentDigests": item["variantContentDigests"],
            "hasConflictingInventoryVariants": item["hasConflictingVariants"],
            "publisherSourceRecord": source_record or NOT_STORED,
            "executionContract": contract,
            "censusFaceKeys": [],
        }

    face_records = []
    binding_state_counts = Counter()
    unresolved_records = []
    for source_face in census_rows:
        runtime_ids, match_method, match_evidence = match_contracts(source_face, inventory)
        for runtime_id in runtime_ids:
            runtime_bindings[runtime_id]["censusFaceKeys"].append(source_face["censusFaceKey"])
        face_ids = sorted({
            canonical_face_id(inventory[runtime_id]["package"])
            for runtime_id in runtime_ids
            if canonical_face_id(inventory[runtime_id]["package"])
        })
        if not runtime_ids:
            binding_state = "NO_REGISTERED_RUNTIME_BINDING"
            resolved_face_id = UNKNOWN
        elif len(face_ids) == 1:
            binding_state = "ONE_REGISTERED_CANONICAL_FACE"
            resolved_face_id = face_ids[0]
        else:
            binding_state = "MULTIPLE_REGISTERED_CANONICAL_FACES"
            resolved_face_id = UNKNOWN
        binding_state_counts[binding_state] += 1

        summaries = [runtime_bindings[runtime_id]["summary"] for runtime_id in runtime_ids]
        qt_identity = source_face["inventoryRecordIdsOrQtIdentityId"] if source_face["printer"] == "Qualification Targets" else ""
        qt_record = qt_chairs.get(qt_identity)
        history_records = []
        for runtime_id in runtime_ids:
            history_records.extend(founder_by_runtime.get(runtime_id, []))
        for face_id in face_ids:
            history_records.extend(founder_by_face.get(face_id, []))
        dedup_history = {canonical_json(record): record for record in history_records}
        workbook_sha = source_face["workbookAssetSha256"]
        sha_consistency = (
            "NOT_APPLICABLE_NO_WORKBOOK_SHA"
            if workbook_sha == NOT_STORED
            else ("MATCH" if workbook_sha in face_ids else "MISMATCH_OR_NO_BINDING")
        )

        face_record = {
            "identity": {
                "censusFaceKey": source_face["censusFaceKey"],
                "masterNumber": source_face["masterNumber"],
                "canonicalFaceId": resolved_face_id,
                "candidateCanonicalFaceIds": face_ids,
                "canonicalFaceResolution": binding_state,
                "printer": source_face["printer"],
                "targetName": source_face["targetName"],
                "sourceRecordCount": source_face["sourceRecordCount"],
                "inventoryNumbersOrQtOrdinal": source_face["inventoryNumbersOrQtOrdinal"],
                "inventoryRecordIdsOrQtIdentityId": source_face["inventoryRecordIdsOrQtIdentityId"],
                "printerProductIdsOrSkus": source_face["printerProductIdsOrSkus"],
                "runtimeIds": runtime_ids,
                "aliasesAndAlternateIdentities": [
                    {
                        "runtimeId": summary["runtimeId"],
                        "productId": summary["productId"],
                        "productName": summary["productName"],
                        "sku": summary["sku"],
                    }
                    for summary in summaries
                ],
            },
            "existingUnderstanding": {
                "assignedFamilyOrQtFamily": source_face["assignedFamilyOrQtFamily"],
                "smartFamily": source_face["smartFamily"],
                "statusOrQtCapability": source_face["statusOrQtCapability"],
                "explicitGovernedLabels": source_face["explicitGovernedLabels"],
                "explicitGovernedValues": source_face["explicitGovernedValues"],
                "registeredSemanticRegionsOrObjects": source_face["registeredSemanticRegionsOrObjects"],
                "classificationBasis": source_face["classificationBasis"],
                "backendMissionFamilies": distinct_non_missing([summary["missionFamily"] for summary in summaries]),
                "backendResultPackageTypes": distinct_non_missing([summary["resultPackageType"] for summary in summaries]),
                "backendGeometryLayoutModels": distinct_non_missing([summary["geometryLayoutModel"] for summary in summaries]),
                "backendScoringModels": distinct_non_missing([summary["scoringModel"] for summary in summaries]),
                "qtBackendAuthorityRecord": qt_record or NOT_STORED,
            },
            "authority": {
                "recordSource": source_face["recordSource"],
                "sourceUrlsOrPublisherEvidence": source_face["sourceUrlsOrPublisherEvidence"],
                "workbookImage": source_face["image"],
                "workbookAssetSha256": workbook_sha,
                "registeredAssetShaConsistency": sha_consistency,
                "storedGoverningRule": source_face["storedGoverningRule"],
                "authorityLimitOrUnresolved": source_face["authorityLimitOrUnresolved"],
                "forbiddenInference": source_face["forbiddenInference"],
                "contractAuthorityVersions": distinct_non_missing([summary["authorityContractVersion"] for summary in summaries]),
                "contractPublisherRuleAuthorities": distinct_non_missing([summary["publisherRuleAuthority"] for summary in summaries]),
                "truthLayerProvenance": {
                    "publisherOrGoverningTruth": [
                        "authority.sourceUrlsOrPublisherEvidence",
                        "authority.storedGoverningRule",
                        "authority.contractPublisherRuleAuthorities",
                    ],
                    "physicalOrFaceTruth": [
                        "authority.workbookAssetSha256",
                        "identity.candidateCanonicalFaceIds",
                        "runtimeBindings[*].executionContract.components.registration",
                        "runtimeBindings[*].executionContract.components.geometry",
                    ],
                    "sczn3DerivedIntelligence": [
                        "runtimeBindings[*].executionContract.components.scoring",
                        "runtimeBindings[*].executionContract.components.mission",
                        "history.founderRecords",
                    ],
                    "machineOrRuntimeState": [
                        "executionState",
                        "runtimeBindings[*].executionContract",
                    ],
                },
            },
            "reusableMachinery": {
                "missionFamilies": distinct_non_missing([summary["missionFamily"] for summary in summaries]),
                "resultPackageTypes": distinct_non_missing([summary["resultPackageType"] for summary in summaries]),
                "executors": distinct_non_missing([summary["executor"] for summary in summaries]),
                "secAdapters": distinct_non_missing([summary["secAdapter"] for summary in summaries]),
                "preservationValidators": distinct_non_missing([summary["preservationValidator"] for summary in summaries]),
                "scoringModels": distinct_non_missing([summary["scoringModel"] for summary in summaries]),
                "classificationFootprintModes": distinct_non_missing([summary["classificationFootprintMode"] for summary in summaries]),
                "outsideRules": distinct_non_missing([summary["outsideRule"] for summary in summaries]),
                "capabilityStates": distinct_non_missing([summary["capabilityStates"] for summary in summaries]),
                "requiredPersistenceFromCensus": source_face["requiredPersistence"],
            },
            "history": {
                "founderRecords": list(dedup_history.values()) or NOT_STORED,
                "priorRepairsOrInheritedLessons": founder_lessons(list(dedup_history.values())),
                "exceptionsOrHoldsFromCensus": source_face["authorityLimitOrUnresolved"],
            },
            "executionState": {
                "runtimeBindingMatchMethod": match_method,
                "runtimeBindingMatchEvidence": match_evidence,
                "contractLifecycleStatuses": distinct_non_missing([summary["lifecycleStatus"] for summary in summaries]),
                "contractSupportStatuses": distinct_non_missing([summary["supportStatus"] for summary in summaries]),
                "contractApprovals": distinct_non_missing([summary["approval"] for summary in summaries]),
                "geometryStatuses": distinct_non_missing([summary["geometryStatus"] for summary in summaries]),
                "scoringStatuses": distinct_non_missing([summary["scoringStatus"] for summary in summaries]),
                "knownContractVariantConflict": any(runtime_bindings[rid]["hasConflictingInventoryVariants"] for rid in runtime_ids),
            },
            "provenance": {
                "censusWorkbook": relative(CENSUS),
                "censusSheet": CENSUS_SHEET,
                "censusRow": source_face["masterNumber"] + 1,
                "runtimeBindingReferences": runtime_ids,
                "qtAuthorityReference": relative(QT_AUTHORITY) if qt_record else NOT_STORED,
            },
        }
        face_records.append(face_record)

        if binding_state != "ONE_REGISTERED_CANONICAL_FACE" or sha_consistency == "MISMATCH_OR_NO_BINDING":
            reasons = []
            if binding_state == "NO_REGISTERED_RUNTIME_BINDING":
                reasons.append("NO_REGISTERED_RUNTIME_BINDING")
            if binding_state == "MULTIPLE_REGISTERED_CANONICAL_FACES":
                reasons.append("MULTIPLE_REGISTERED_CANONICAL_FACES")
            if sha_consistency == "MISMATCH_OR_NO_BINDING":
                reasons.append("WORKBOOK_SHA_DOES_NOT_RECONCILE_TO_REGISTERED_FACE")
            unresolved_records.append({
                "censusFaceKey": source_face["censusFaceKey"],
                "masterNumber": source_face["masterNumber"],
                "printer": source_face["printer"],
                "targetName": source_face["targetName"],
                "reasons": reasons,
                "runtimeIds": runtime_ids,
                "candidateCanonicalFaceIds": face_ids,
                "matchMethod": match_method,
            })

    canonical_groups: dict[str, dict] = {}
    for runtime_id, binding in runtime_bindings.items():
        face_id = binding["summary"]["canonicalFaceId"]
        group = canonical_groups.setdefault(face_id, {
            "canonicalFaceId": face_id,
            "runtimeIds": [],
            "censusFaceKeys": [],
            "printers": [],
            "productNames": [],
            "skus": [],
        })
        group["runtimeIds"].append(runtime_id)
        group["censusFaceKeys"].extend(binding["censusFaceKeys"])
        group["printers"].append(binding["summary"]["printer"])
        group["productNames"].append(binding["summary"]["productName"])
        group["skus"].append(binding["summary"]["sku"])
    for group in canonical_groups.values():
        for key in ("runtimeIds", "censusFaceKeys", "printers", "productNames", "skus"):
            group[key] = sorted(set(group[key]))
        group["runtimeAliasCount"] = len(group["runtimeIds"])
        group["censusReferenceCount"] = len(group["censusFaceKeys"])

    observed_totals = Counter(record["identity"]["printer"] for record in face_records)
    accounted_keys = [record["identity"]["censusFaceKey"] for record in face_records]
    if len(face_records) != EXPECTED_TOTAL or len(accounted_keys) != len(set(accounted_keys)):
        raise RuntimeError("authoritative census row accounting failed")
    if dict(observed_totals) != EXPECTED_PRINTER_TOTALS:
        raise RuntimeError(f"printer totals do not reconcile: {dict(observed_totals)}")
    if workbook_totals != EXPECTED_PRINTER_TOTALS:
        raise RuntimeError(f"workbook printer totals changed: {workbook_totals}")

    unmatched_runtime = sorted(runtime_id for runtime_id, item in runtime_bindings.items() if not item["censusFaceKeys"])
    issue_counts = Counter(reason for record in unresolved_records for reason in record["reasons"])
    summary = {
        "schemaVersion": "sczn3-episode64-inheritance-census-accounting-v1",
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "sourceCommit": "174d2d07297bff711b058619c9f1527049e087db",
        "mechanicalClosure": {"input": EXPECTED_TOTAL, "accounted": len(face_records), "unaccounted": 0, "status": "PASS"},
        "printerTotals": {
            printer: {"expected": total, "observed": observed_totals[printer], "status": "PASS"}
            for printer, total in EXPECTED_PRINTER_TOTALS.items()
        },
        "smartFamilyTotals": dict(sorted(Counter(record["existingUnderstanding"]["smartFamily"] for record in face_records).items())),
        "mechanicalRuntimeBindingStates": dict(sorted(binding_state_counts.items())),
        "registeredCanonicalFaceIndex": {
            "uniqueFaceIdsInBackend": len(canonical_groups),
            "uniqueFaceIdsReferencedByCensus": len({face_id for record in face_records for face_id in record["identity"]["candidateCanonicalFaceIds"]}),
            "backendFaceGroupsWithRuntimeAliases": sum(1 for group in canonical_groups.values() if group["runtimeAliasCount"] > 1),
            "backendFaceGroupsReferencedByMultipleCensusRows": sum(1 for group in canonical_groups.values() if group["censusReferenceCount"] > 1),
        },
        "runtimeContractInventory": {
            "selectedCanonicalRuntimeBindings": len(runtime_bindings),
            "contractIdentitiesWithConflictingFileVariants": len(contract_conflicts),
            "runtimeBindingsNotJoinedToAuthoritativeCensus": len(unmatched_runtime),
        },
        "identityOrBackendStateReconciliationIssues": {
            "censusRecordsListed": len(unresolved_records),
            "issueCounts": dict(sorted(issue_counts.items())),
            "contractVariantConflictCount": len(contract_conflicts),
            "unmatchedRuntimeBindingCount": len(unmatched_runtime),
        },
        "semanticClassificationPerformed": False,
        "prohibitedSemanticBucketsAbsent": ["Full", "Partial", "New", "Authority Blocked", "Identity Blocked"],
    }

    full = {
        "schemaVersion": "sczn3-episode64-the13-inheritance-census-input-v1",
        "generatedAt": summary["generatedAt"],
        "operation": "read_only_backend_knowledge_export",
        "sourceCommit": summary["sourceCommit"],
        "criticalRule": "No missing semantic intelligence was inferred. UNKNOWN and NOT_STORED are preserved as data.",
        "mechanicalClosure": summary["mechanicalClosure"],
        "globalAuthorityReferences": {
            "printerShooterSczn3Authority": relative(PRINTER_SHOOTER_AUTHORITY),
            "x10OrderedGrammarAuthority": relative(X10_AUTHORITY),
            "scoreOnTouchAuthority": {"source": relative(SCORE_ON_TOUCH), "record": read_json(SCORE_ON_TOUCH)},
            "founder20Of20Closure": {"source": relative(FOUNDER_CLOSURE), "record": founder_closure},
            "qtAuthorityHeader": qt_authority_header,
        },
        "faces": face_records,
        "runtimeBindings": runtime_bindings,
        "canonicalFaceGroups": canonical_groups,
    }

    unresolved = {
        "schemaVersion": "sczn3-episode64-inheritance-census-unreconciled-v1",
        "generatedAt": summary["generatedAt"],
        "note": "These are mechanical identity/runtime-state issues, not LLM inheritance classifications.",
        "censusRecords": unresolved_records,
        "contractFileVariantConflicts": contract_conflicts,
        "runtimeBindingsNotJoinedToAuthoritativeCensus": unmatched_runtime,
    }

    schema = {
        "schemaVersion": "sczn3-episode64-inheritance-census-data-dictionary-v1",
        "rootCollections": {
            "faces": "Exactly 1,993 authoritative workbook rows, one per census face key.",
            "runtimeBindings": "Current canonical-filename Backend execution contracts, normalized by runtime ID. Complete stored contracts are included without semantic rewriting.",
            "canonicalFaceGroups": "Runtime aliases grouped only by stored registration assetSha256.",
        },
        "faceSections": {
            "identity": "Census identity, registered face resolution, aliases, product/SKU and runtime bindings.",
            "existingUnderstanding": "Only classifications, structures, and treatment data already stored in the workbook, QT authority, or linked contracts.",
            "authority": "Publisher/governing, physical/face, SCZN3-derived, and runtime provenance boundaries.",
            "reusableMachinery": "Stored executor/result/scoring/footprint/persistence machinery associated with linked runtime contracts.",
            "history": "Stored Founder decisions, repairs, and approval evidence tied by runtime ID or canonical face hash.",
            "executionState": "Mechanical contract/binding/lifecycle state. It is not a semantic inheritance grade.",
            "provenance": "Exact source references for the face record.",
        },
        "sentinels": {"UNKNOWN": "A required identity/result cannot be resolved from stored evidence.", "NOT_STORED": "The source has no stored value for the field."},
        "canonicalFaceResolutionValues": {
            "ONE_REGISTERED_CANONICAL_FACE": "All matched runtime identities share one stored registration asset SHA-256.",
            "MULTIPLE_REGISTERED_CANONICAL_FACES": "The census row mechanically matches runtime identities with more than one registration asset SHA-256; no face was selected.",
            "NO_REGISTERED_RUNTIME_BINDING": "No exact QT identity, SKU/product token, or exact normalized product-name binding was found.",
        },
        "matchMethods": {
            "exact_qt_identity_id": "QT catalog/identity ID equality.",
            "exact_sku_or_product_id_token": "Exact normalized token equality against stored SKU/product ID.",
            "exact_normalized_product_name": "Exact normalized product-name equality, used only when no token match exists.",
            "NO_RUNTIME_BINDING_MATCH": "No exact mechanical join succeeded.",
        },
        "nonGoals": [
            "No Full/Partial/New/Authority Blocked/Identity Blocked classification.",
            "No target implementation or repair.",
            "No contract mutation.",
            "No inference of missing regions, rules, values, or authority.",
        ],
    }

    provenance_sources = [
        (CENSUS, "authoritative 1,993-face census and printer totals"),
        (PUBLISHER_SOURCE, "publisher source records for printers 1-12"),
        (QT_AUTHORITY, "stored Qualification Targets authority and treatments"),
        (FOUNDER_LEDGER, "Founder walk history for early entries"),
        (FOUNDER_CLOSURE, "Founder 20/20 unique canonical-face closure"),
        (SCORE_ON_TOUCH, "shared Score on Touch Founder authority"),
        (PRINTER_SHOOTER_AUTHORITY, "Founder printer/shooter/SCZN3 authority boundary"),
        (X10_AUTHORITY, "Founder X=10 ordered-grammar authority"),
    ]
    provenance = {
        "schemaVersion": "sczn3-episode64-inheritance-census-provenance-v1",
        "generatedAt": summary["generatedAt"],
        "sourceCommit": summary["sourceCommit"],
        "sources": [
            {"path": relative(path), "sha256": sha256(path), "bytes": path.stat().st_size, "role": role}
            for path, role in provenance_sources
        ],
        "contractInventory": {
            "directory": relative(CONTRACTS),
            "jsonFileCountIncludingVariants": len(list(CONTRACTS.glob("*.json"))),
            "selectedCanonicalRuntimeBindings": len(runtime_bindings),
            "selectionRule": "prefer the contract filename without a trailing ' N' copy suffix; retain every variant path/digest and report conflicting content",
        },
        "joinRules": schema["matchMethods"],
        "outputNormalization": "Full execution contracts are stored once in runtimeBindings; face records reference runtime IDs and contain compact exact aggregates.",
    }

    report = "\n".join([
        "# Episode 64 — THE 13 Inheritance-Census Input",
        "",
        f"Source commit: `{summary['sourceCommit']}`",
        "",
        "## Mechanical closure",
        "",
        f"- INPUT: **{summary['mechanicalClosure']['input']:,}**",
        f"- ACCOUNTED: **{summary['mechanicalClosure']['accounted']:,}**",
        f"- UNACCOUNTED: **{summary['mechanicalClosure']['unaccounted']:,}**",
        f"- Status: **{summary['mechanicalClosure']['status']}**",
        "",
        "All 13 printer totals match the authoritative workbook.",
        "",
        "## Mechanical runtime-binding state",
        "",
        *[f"- {key}: **{value:,}**" for key, value in sorted(binding_state_counts.items())],
        "",
        "These counts are not semantic inheritance classifications. The LLM must determine legitimate inheritance separately.",
        "",
        "## Reconciliation evidence",
        "",
        f"- Registered Backend runtime bindings indexed: **{len(runtime_bindings):,}**",
        f"- Registered Backend canonical face hashes indexed: **{len(canonical_groups):,}**",
        f"- Census records requiring explicit identity/runtime review: **{len(unresolved_records):,}**",
        f"- Contract identities with conflicting file variants preserved and reported: **{len(contract_conflicts):,}**",
        f"- Runtime bindings not joined to a 1,993 census row: **{len(unmatched_runtime):,}**",
        "",
        "No missing target semantics were inferred, no target contract was edited, and no production cohort was started.",
        "",
    ])
    return full, summary, schema, {"unreconciled": unresolved, "provenance": provenance}, report


def write_json(path: Path, payload: Any, *, compact: bool = False) -> None:
    if compact:
        text = json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n"
    else:
        text = json.dumps(payload, indent=2, ensure_ascii=False) + "\n"
    path.write_text(text, encoding="utf-8")


def main() -> None:
    full, summary, schema, sidecars, report = build()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    write_json(FULL_CENSUS, full, compact=True)
    write_json(SUMMARY, summary)
    write_json(SCHEMA, schema)
    write_json(PROVENANCE, sidecars["provenance"])
    write_json(UNRECONCILED, sidecars["unreconciled"])
    REPORT.write_text(report, encoding="utf-8")
    payload_outputs = [FULL_CENSUS, SUMMARY, SCHEMA, PROVENANCE, UNRECONCILED, REPORT]
    delivery = {
        "schemaVersion": "sczn3-episode64-inheritance-census-delivery-v1",
        "sourceCommit": summary["sourceCommit"],
        "mechanicalClosure": summary["mechanicalClosure"],
        "artifacts": [
            {"path": relative(path), "bytes": path.stat().st_size, "sha256": sha256(path)}
            for path in payload_outputs
        ],
    }
    write_json(DELIVERY, delivery)
    outputs = [*payload_outputs, DELIVERY]
    print(json.dumps({
        "status": "PASS",
        "mechanicalClosure": summary["mechanicalClosure"],
        "outputs": [{"path": relative(path), "bytes": path.stat().st_size, "sha256": sha256(path)} for path in outputs],
    }, indent=2))


if __name__ == "__main__":
    main()
